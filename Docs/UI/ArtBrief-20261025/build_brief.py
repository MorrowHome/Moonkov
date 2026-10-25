"""Build the paired functional briefs with python-docx from brief-content.json.

Use the Python runtime returned by the Codex workspace dependency loader.
Render outputs with the documents skill renderer before delivery.
"""

from pathlib import Path
import json

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image


ROOT = Path(__file__).resolve().parent
DATA = json.loads((ROOT / "brief-content.json").read_text(encoding="utf-8"))


def set_font(run, lang, size=None, bold=None):
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    lng = OxmlElement("w:lang")
    lng.set(qn("w:val"), "zh-CN" if lang == "zh" else "en-US")
    lng.set(qn("w:eastAsia"), "zh-CN")
    rpr.append(lng)


def add_text(doc, text, lang, style=None, size=None, bold=None):
    para = doc.add_paragraph(style=style)
    fmt = para.paragraph_format
    fmt.line_spacing = Pt(17)
    fmt.space_after = Pt(6)
    if style in ("Title", "Heading 1", "Heading 2"):
        fmt.line_spacing = Pt(28 if style == "Title" else 24 if style == "Heading 1" else 20)
        fmt.space_before = Pt(0 if style in ("Title", "Heading 1") else 12)
        fmt.space_after = Pt(8 if style in ("Title", "Heading 1") else 5)
        fmt.keep_with_next = True
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    para._p.get_or_add_pPr().append(snap)
    set_font(para.add_run(text), lang, size, bold)
    return para


def prevent_row_split(row):
    prop = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:cantSplit")
    prop.append(flag)


def add_table(doc, block, lang):
    table = doc.add_table(rows=1, cols=len(block["headers"]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = block.get("widths", [6.9 / len(block["headers"])] * len(block["headers"]))
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:color"), "D9D9D9")
        borders.append(el)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", 85), ("bottom", 85), ("left", 100), ("right", 100)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    props.append(margins)
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    all_rows = [block["headers"], *block["rows"]]
    for ri, values in enumerate(all_rows):
        row = table.rows[0] if ri == 0 else table.add_row()
        prevent_row_split(row)
        for ci, value in enumerate(values):
            cell = row.cells[ci]
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = Pt(15.5)
            snap = OxmlElement("w:snapToGrid")
            snap.set(qn("w:val"), "0")
            para._p.get_or_add_pPr().append(snap)
            if ci == 0 and len(values) == 3:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(para.add_run(value[lang]), lang, 10.5, ri == 0)
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E7EDF2")
                cell._tc.get_or_add_tcPr().append(shade)
        if ri == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(1)
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.line_spacing = 0.4
    gap.add_run().font.size = Pt(2)


def add_screenshot(doc, block, lang):
    """Keep a UI name, implementation image and short caption on one page."""
    image_path = ROOT / block["image"]
    if not image_path.is_file():
        raise FileNotFoundError(f"Screenshot not found: {image_path}")

    title = add_text(doc, block["title"][lang], lang, size=12.5, bold=True)
    title.paragraph_format.line_spacing = Pt(17)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    title.paragraph_format.keep_with_next = True
    title.paragraph_format.keep_together = True

    text = block.get("text", {}).get(lang, "")
    if text:
        description = add_text(doc, text, lang, size=10)
        description.paragraph_format.line_spacing = Pt(13)
        description.paragraph_format.space_after = Pt(3)
        description.paragraph_format.keep_with_next = True
        description.paragraph_format.keep_together = True

    with Image.open(image_path) as source:
        pixel_width, pixel_height = source.size
    if pixel_width <= 0 or pixel_height <= 0:
        raise ValueError(f"Screenshot has invalid dimensions: {image_path}")
    width = min(6.9, 3.75 * pixel_width / pixel_height)
    height = width * pixel_height / pixel_width

    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.line_spacing = 1
    picture.paragraph_format.space_before = Pt(0)
    picture.paragraph_format.space_after = Pt(2)
    picture.paragraph_format.keep_with_next = True
    picture.paragraph_format.keep_together = True
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    picture._p.get_or_add_pPr().append(snap)
    shape = picture.add_run().add_picture(str(image_path), width=Inches(width), height=Inches(height))
    shape._inline.docPr.set("descr", block["title"][lang])

    note = add_text(doc, block.get("note", {}).get(lang, ""), lang, size=9.5)
    note.paragraph_format.line_spacing = Pt(12)
    note.paragraph_format.space_before = Pt(0)
    note.paragraph_format.space_after = Pt(8)
    note.paragraph_format.keep_with_next = False
    note.paragraph_format.keep_together = True


def build(lang):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    sec.header_distance = Inches(0.25)
    sec.footer_distance = Inches(0.25)

    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.paragraph_format.widow_control = True
    # python-docx's default template contains a blue title border. Remove it.
    for element in list(doc.styles.element.iter()):
        if element.tag == qn("w:pBdr"):
            element.getparent().remove(element)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(6)
    for name, size, before, after in (("Title", 22, 0, 10), ("Heading 1", 18, 0, 10), ("Heading 2", 13, 10, 5)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = sec.header.paragraphs[0]
    set_font(header.add_run("MOONKOV  /  UI MODULES"), lang, 9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("2026-10-25  |  "), lang, 9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    for page_index, page in enumerate(DATA["pages"]):
        if page_index:
            doc.add_page_break()
        if page.get("title"):
            add_text(doc, page["title"][lang], lang, "Title" if page_index == 0 else "Heading 1")
        for block in page["blocks"]:
            if block["type"] == "p":
                add_text(doc, block["text"][lang], lang)
            elif block["type"] == "h":
                add_text(doc, block["text"][lang], lang, "Heading 2")
            elif block["type"] == "table":
                add_table(doc, block, lang)
            elif block["type"] == "screenshot":
                add_screenshot(doc, block, lang)
            elif block["type"] == "bullets":
                for item in block["items"]:
                    add_text(doc, item[lang], lang, "List Bullet")
            elif block["type"] == "outline":
                for group in block["groups"]:
                    add_text(doc, group["title"][lang], lang, "Heading 2")
                    for item in group["items"]:
                        para = add_text(doc, item[lang], lang)
                        para.paragraph_format.left_indent = Inches(0.2)
            elif block["type"] == "nine_slice":
                labels = block["labels"][lang]
                block_table = {
                    "headers": [{lang: labels[i]} for i in range(3)],
                    "rows": [
                        [{lang: labels[i]} for i in range(3, 6)],
                        [{lang: labels[i]} for i in range(6, 9)],
                    ],
                    "widths": [1.4, 4.1, 1.4],
                }
                add_table(doc, block_table, lang)

    document_title = DATA.get("title") or DATA["pages"][0].get("title") or {
        "zh": "Moonkov UI 图册", "en": "Moonkov UI Screenshots"
    }
    doc.core_properties.title = document_title[lang]
    doc.core_properties.subject = "Functional modules and hierarchy for UI art planning"
    doc.core_properties.author = "Moonkov"
    doc.core_properties.keywords = "Moonkov UI modules hierarchy art brief"
    filename = "Moonkov_UI_Modules_zh-CN.docx" if lang == "zh" else "Moonkov_UI_Modules_en.docx"
    output = ROOT / filename
    doc.save(output)
    print(output)


if __name__ == "__main__":
    for language in DATA["languages"]:
        build(language)
